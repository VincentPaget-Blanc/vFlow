"""Reproduced RC4 regressions: numeric truth, relational integrity and protected files."""
import copy
import gzip
import json
from dataclasses import asdict
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from tests.test_statistical_audit import samples,hierarchical
from vflow.statistics import AuditSample,run_audit
from vflow.statistics.wasserstein import sorted_w1
from vflow.statistics.audit_serialization import validate_bundle,commit_bundle,append_decision,read_bundle
from vflow.statistics.audit_export import export_xlsx,export_csv_package
from vflow.workspace.model import Workspace,FileReference,SampleReference

def tiny_hierarchy(magnitude):
    s=[AuditSample(str(i),str(i),pd.DataFrame({'x':[-magnitude,0.,magnitude]})) for i in range(6)]
    return s,{str(i):str(i//2) for i in range(6)}

def test_nonconstant_within_sample_variance_cannot_underflow_to_identical():
    s,m=tiny_hierarchy(1e-200)
    with pytest.raises(ValueError,match='numeric range'):
        run_audit(s,audit_type='hierarchy',hierarchy_mapping=m,policy={'n_projections':2})

def test_representable_extremely_small_variance_matches_residual_truth():
    s,m=tiny_hierarchy(1e-150)
    result=run_audit(s,audit_type='hierarchy',hierarchy_mapping=m,policy={'n_projections':2})
    model=result['variance_components'][0]
    assert model['model_status']=='boundary'
    assert model['sigma2_particle']/1e-300==pytest.approx(12/17,rel=1e-4)
    assert model['icc_bio']==0

def test_subnormal_positive_raw_distance_is_not_saved_as_zero():
    step=np.nextafter(0.,1.)
    with pytest.raises(ValueError,match='numeric range'):
        sorted_w1(np.array([0.,step,2*step]),np.array([0.,step,3*step]))

def test_unidentifiable_model_can_still_report_finite_large_biological_center():
    s=[AuditSample(str(i),str(i),pd.DataFrame({'x':[8e307,9e307,1e308]})) for i in range(2)]
    result=run_audit(s,audit_type='hierarchy',hierarchy_mapping={'0':'same','1':'same'},policy={'n_projections':2})
    assert result['variance_components'][0]['model_status']=='unidentifiable'
    assert result['equal_bio_centers'][0]['equal_bio_center']==pytest.approx(9e307)
    assert result['influence'][0]['standardized_change_condition']=='insufficient_biological_units'

@pytest.mark.parametrize('defect',['asymmetric_matrix','diagonal','missing_joint_distance','overcount','variable_identity',
    'missing_description','duplicate_variable','matrix_variables','population_hash','policy','distance_disagreement'])
def test_relationally_invalid_bundle_is_rejected(defect):
    b=run_audit(samples(n=3),policy={'n_projections':2})
    if defect=='asymmetric_matrix':b['global_distance_matrix'][0][1]+=1
    elif defect=='diagonal':b['global_distance_matrix'][0][0]=1
    elif defect=='missing_joint_distance':b['global_distance_matrix'][0][1]=b['global_distance_matrix'][1][0]=None
    elif defect=='overcount':b['sample_qc'][0]['multivariate_complete_count']=10000
    elif defect=='variable_identity':b['variable_qc'][0]['sample_id']='unknown'
    elif defect=='missing_description':b['variable_qc'][0].pop('median')
    elif defect=='duplicate_variable':b['variable_qc'].append(b['variable_qc'][0].copy())
    elif defect=='matrix_variables':b['variable_matrices'].pop(b['variables'][0])
    elif defect=='population_hash':b['samples'][0]['population_snapshot']={'gate':'changed'}
    elif defect=='policy':b['policy']['isolation_threshold']=float('nan')
    elif defect=='distance_disagreement':b['variable_distances'][0]['normalized_w1']+=1
    with pytest.raises(ValueError):validate_bundle(b)

def test_bundle_reference_selection_must_match_payload(tmp_path):
    b=run_audit(samples(n=2),policy={'n_projections':2});ref=commit_bundle(Workspace(),b,state_dir=tmp_path)
    ref.selected_sample_ids=['another sample']
    with pytest.raises(ValueError,match='selection'):read_bundle(ref)

@pytest.mark.parametrize('exporter',[export_xlsx,export_csv_package])
@pytest.mark.parametrize('kind',['source','audit','gate','workspace'])
def test_export_cannot_overwrite_project_or_acquisition_files(tmp_path,exporter,kind):
    s=samples(n=2);w=Workspace()
    for sample in s:
        path=tmp_path/(sample.sample_id+'.csv');sample.dataframe.to_csv(path,index=False)
        source=FileReference.create(path);sample.source=asdict(source);w.samples[sample.sample_id]=SampleReference(sample.sample_id,sample.display_name,source)
    b=run_audit(s,policy={'n_projections':2});ref=commit_bundle(w,b,state_dir=tmp_path)
    gate=tmp_path/'gate.json';gate.write_text('{"gate":"original"}')
    w.default_gate_session=FileReference.create(gate,kind='gates')
    workspace_path=tmp_path/'project.vflow';w.save(workspace_path)
    target={'source':tmp_path/'0.csv','audit':Path(ref.result.absolute_path),'gate':gate,'workspace':workspace_path}[kind]
    before=target.read_bytes()
    with pytest.raises(ValueError,match='cannot replace'):
        exporter(b,target,ref,w,protected_paths=[workspace_path])
    assert target.read_bytes()==before

def test_source_symlink_export_cannot_bypass_protection(tmp_path):
    source=tmp_path/'source.csv';source.write_text('x\n1\n2\n3\n')
    b=run_audit(samples(n=1),policy={'n_projections':2});b['samples'][0]['source']=asdict(FileReference.create(source))
    alias=tmp_path/'alias.xlsx';alias.symlink_to(source)
    with pytest.raises(ValueError,match='cannot replace'):export_xlsx(b,alias)
    assert alias.is_symlink() and source.read_text()=='x\n1\n2\n3\n'

def test_excel_forbidden_unicode_is_escaped_and_review_json_remains_exact(tmp_path):
    b=run_audit(samples(n=1),policy={'n_projections':2});w=Workspace();ref=commit_bundle(w,b,state_dir=tmp_path)
    note='review \uffff and \ufffe and emoji \U0001f600';append_decision(ref,'0','note',note=note)
    target=tmp_path/'review.xlsx';export_xlsx(b,target,ref,w)
    book=load_workbook(target);chunks=[r[1].value for r in book['Provenance'] if str(r[0].value).startswith('review_context_chunk_')]
    assert json.loads(''.join(chunks))['decision_history'][0]['note']==note
    assert '\\uffff' in book['Flags_Decisions'].cell(3,6).value

def test_workspace_save_failure_does_not_rebase_live_references(tmp_path,monkeypatch):
    import vflow.workspace.model as module
    source=tmp_path/'source.csv';source.write_text('x\n1\n')
    original=tmp_path/'old/project.vflow';w=Workspace();ref=FileReference.create(source,original)
    w.samples['s']=SampleReference('s','source',ref);before=copy.deepcopy(w.payload())
    monkeypatch.setattr(module,'atomic_json',lambda *a,**k:(_ for _ in ()).throw(OSError('disk full')))
    with pytest.raises(OSError):w.save(tmp_path/'elsewhere/new.vflow')
    assert w.payload()==before and w.samples['s'].source is ref

@pytest.mark.parametrize('defect',['audit_container','unsafe_id','selection','review_history','source_path'])
def test_malformed_workspace_audit_metadata_is_rejected(tmp_path,defect):
    b=run_audit(samples(n=1),policy={'n_projections':2});w=Workspace();ref=commit_bundle(w,b,state_dir=tmp_path)
    data=w.payload();entry=data['audits'][ref.audit_id]
    if defect=='audit_container':data['audits']=[]
    elif defect=='unsafe_id':entry['audit_id']='../../outside'
    elif defect=='selection':entry['selected_sample_ids']='0'
    elif defect=='review_history':entry['decision_history']=[{'action':'exclude','sample_id':'missing'}]
    elif defect=='source_path':entry['result']['absolute_path']=42
    path=tmp_path/'broken.vflow';path.write_text(json.dumps(data))
    with pytest.raises(ValueError):Workspace.load(path)

def test_qc_rejects_hierarchy_metadata_that_would_mislabel_summary():
    with pytest.raises(ValueError,match='hierarchical'):
        run_audit(samples(n=1),hierarchy_mapping={'0':'bio'})

@pytest.mark.parametrize('input_case',['missing_samples','blank_id','blank_name','numeric_name','numeric_population','numeric_column'])
def test_invalid_public_inputs_fail_before_worker_calculations(input_case):
    s=samples(n=1);kwargs={}
    if input_case=='missing_samples':s=None
    elif input_case=='blank_id':s[0].sample_id=' '
    elif input_case=='blank_name':s[0].display_name=' '
    elif input_case=='numeric_name':kwargs['name']=42
    elif input_case=='numeric_population':kwargs['population_label']=42
    elif input_case=='numeric_column':s[0].dataframe.columns=[1,2,3]
    with pytest.raises(ValueError):run_audit(s,**kwargs)

@pytest.mark.parametrize('scenario',['tiny_sd','large_mean','tiny_w1'])
def test_numeric_handling_when_longdouble_is_only_float64(monkeypatch,scenario):
    from vflow.statistics.variable_eligibility import describe,robust_scale
    monkeypatch.setattr(np,'longdouble',np.float64)
    if scenario=='tiny_sd':
        row=describe([-1e-200,0.,1e-200]);assert row['std']==pytest.approx(1e-200,rel=1e-12,abs=0)
        assert robust_scale([0,0,0,0,1e-200])[1]==pytest.approx(4e-201,rel=1e-12,abs=0)
    elif scenario=='large_mean':
        row=describe([8e307,9e307,1e308]);assert row['mean']==pytest.approx(9e307)
        assert row['std']==pytest.approx(1e307,rel=1e-12)
        assert robust_scale([0,0,0,0,1e308])[1]==pytest.approx(4e307)
    else:
        step=np.nextafter(0.,1.)
        with pytest.raises(ValueError,match='numeric range'):
            sorted_w1(np.array([0.,step,2*step]),np.array([0.,step,3*step]))
