import copy
import json
import threading
import zipfile
import numpy as np
import pandas as pd
import pytest
from scipy.stats import wasserstein_distance
from vflow.statistics import AuditSample,run_audit,AuditCancelled
from vflow.statistics.variable_eligibility import eligible_variables,robust_scale
from vflow.statistics.wasserstein import sorted_w1
from vflow.statistics.hierarchy import estimate,sufficient_stats
from vflow.statistics.audit_serialization import commit_bundle,read_bundle,append_decision
from vflow.statistics.audit_export import export_xlsx,export_csv_package,SHEETS
from vflow.workspace.model import Workspace,SampleReference,FileReference


def samples(seed=9,n=5,rows=300):
    rng=np.random.default_rng(seed)
    return [AuditSample(str(i),f'Sample {i}',pd.DataFrame(rng.normal(size=(rows+i*10,3)),columns=['Fluorescence','BG2','Distance'])) for i in range(n)]

@pytest.mark.parametrize('n,status',[(1,'insufficient_peers'),(2,'pairwise_only')])
def test_small_n_semantics(n,status):
    result=run_audit(samples(n=n));assert {r['status'] for r in result['sample_qc']}=={status}

def test_concordant_different_counts_no_false_flags_and_all_rows():
    s=samples();r=run_audit(s)
    assert all(x['status']=='concordant' for x in r['sample_qc'])
    assert [x['multivariate_complete_count'] for x in r['sample_qc']]==[len(x.dataframe) for x in s]
    assert not r['event_subsampling']
    mat=np.array(r['global_distance_matrix']);assert np.array_equal(mat,mat.T);assert np.all(np.diag(mat)==0)

def test_shifted_sample_flagged_and_result_deterministic():
    s=samples();s[-1].dataframe+=8
    r=run_audit(s);r2=run_audit(s)
    assert r['sample_qc'][-1]['status']=='flagged_for_review'
    assert r['sample_qc'][-1]['projection_stability']
    assert len(r['sample_qc'][-1]['contributors'])==3
    assert r['global_distance_matrix']==r2['global_distance_matrix']
    assert r['sample_qc']==r2['sample_qc']

def test_low_n_candidate_never_confirmed_outlier():
    s=samples(n=3);s[-1].dataframe+=9;r=run_audit(s)
    assert r['sample_qc'][-1]['status']=='low_n_review_candidate'
    assert r['sample_qc'][-1]['low_n']

def test_joint_only_discordance_with_identical_marginals():
    x=np.linspace(-3,3,500)
    s=[AuditSample(str(i),str(i),pd.DataFrame({'a':x,'b':x if i<4 else x[::-1]})) for i in range(5)]
    r=run_audit(s)
    assert all(row['raw_w1']==0 for row in r['variable_distances'])
    assert r['sample_qc'][-1]['status']=='flagged_for_review'
    assert r['sample_qc'][-1]['projection_stability']

def test_finite_values_and_complete_cases_no_imputation():
    s=samples();s[0].dataframe.loc[:9,'Distance']=np.nan;s[1].dataframe.loc[:4,'BG2']=np.inf
    r=run_audit(s)
    assert r['sample_qc'][0]['multivariate_complete_count']==len(s[0].dataframe)-10
    assert r['sample_qc'][1]['multivariate_complete_count']==len(s[1].dataframe)-5
    assert next(x for x in r['variable_qc'] if x['sample_id']=='0' and x['variable']=='Distance')['missing_count']==10
    assert r['warnings'];json.dumps(r,allow_nan=False)

def test_no_multivariate_complete_rows_mark_unavailable():
    s=samples(rows=30)
    for j,c in enumerate(s[0].dataframe):s[0].dataframe.loc[s[0].dataframe.index%3!=j,c]=np.nan
    r=run_audit(s)
    assert r['sample_qc'][0]['status']=='multivariate_unavailable'
    assert r['sample_qc'][0]['global_peer_distance'] is None
    assert all(row['raw_w1'] is not None for row in r['variable_distances'])

@pytest.mark.parametrize('joint_n,status',[(1,'insufficient_peers'),(2,'pairwise_only'),(3,'low_n_review_candidate'),(4,'low_n_review_candidate')])
def test_missing_joint_samples_cannot_inflate_classification_n(joint_n,status):
    s=samples(n=5,rows=30)
    for sample in s[:5-joint_n]:
        for j,c in enumerate(sample.dataframe):sample.dataframe.loc[sample.dataframe.index%3!=j,c]=np.nan
    if joint_n>=3:s[-1].dataframe+=9
    r=run_audit(s)
    assert r['sample_qc'][-1]['status']==status
    assert r['sample_qc'][-1]['multivariate_sample_count']==joint_n
    assert r['sample_qc'][-1]['global_peer_count']==joint_n-1
    assert r['sample_qc'][-1]['low_n'] and r['sample_qc'][-1]['robust_z'] is None
    assert all(row['status']=='multivariate_unavailable' for row in r['sample_qc'][:5-joint_n])

def review_matrix():
    # A valid symmetric metric: equal-cohesion peers and a moderately separated
    # last sample. The explicit zero-MAD robust-z policy recommends review even
    # though the sample's isolation ratio is below 2.5.
    m=np.ones((7,7));np.fill_diagonal(m,0);m[-1,:-1]=2;m[:-1,-1]=2
    return m

def test_identical_repeat_preserves_robust_z_only_review_candidate():
    from vflow.statistics.discordance import score_samples
    from vflow.statistics.models import POLICY
    m=review_matrix();r=score_samples(samples(n=7),m,{'x':m},m.copy(),POLICY)[-1]
    assert r['global_isolation_ratio']<POLICY['isolation_threshold']
    assert r['robust_z']>=POLICY['robust_z_threshold']
    assert r['status']=='flagged_for_review' and r['projection_stability']
    assert 'robust_z_high' in r['reason_codes'] and 'projection_unstable' not in r['reason_codes']

def test_repeat_disagreement_retains_advisory_unstable_result():
    from vflow.statistics.discordance import score_samples
    from vflow.statistics.models import POLICY
    m=review_matrix();repeat=np.ones_like(m);np.fill_diagonal(repeat,0)
    r=score_samples(samples(n=7),m,{'x':m},repeat,POLICY)[-1]
    assert r['status']=='mildly_discordant' and not r['projection_stability']
    assert 'projection_unstable' in r['reason_codes']

def test_variable_eligibility_and_scale_fallback():
    s=samples();s[0].dataframe['private']=range(len(s[0].dataframe))
    for sample in s:sample.dataframe['label']='a';sample.dataframe['constant']=1
    eligible,excluded,_=eligible_variables(s)
    assert eligible==['Fluorescence','BG2','Distance']
    assert {r['reason'] for r in excluded}=={'missing_in_sample','non_numeric','constant_or_zero_scale'}
    assert robust_scale([0,0,0,0,8])[2]=='pooled_sd'

def test_exact_wide_source_integers_are_not_silently_rounded_for_audit():
    s=samples()
    for a in s:a.dataframe['exact_id']=np.arange(len(a.dataframe),dtype=np.uint64)+np.uint64(2**53+1)
    r=run_audit(s)
    assert 'exact_id' not in r['variables']
    assert {'variable':'exact_id','reason':'integer_precision_exceeds_float64'} in r['excluded_variables']

@pytest.mark.parametrize('a,b',[(np.array([0.,1.]),np.array([1.,2.,3.])),(np.arange(30.),np.linspace(0,8,11))])
def test_w1_against_independent_scipy(a,b):
    assert sorted_w1(np.sort(a),np.sort(b))==pytest.approx(wasserstein_distance(a,b),abs=1e-13)

def test_cancellation_before_any_commit():
    event=threading.Event();event.set()
    with pytest.raises(AuditCancelled):run_audit(samples(),cancel=event)
    stages=[]
    def progress(stage):
        stages.append(stage)
        if 'projection 2/' in stage:event.set()
    event.clear()
    with pytest.raises(AuditCancelled):run_audit(samples(),cancel=event,progress=progress)

def hierarchical(seed=22,bios=20,nested=5,particles=100):
    rng=np.random.default_rng(seed);s=[];mapping={}
    for b in range(bios):
        effect=rng.normal(0,2)
        for j in range(nested):
            sid=f'{b}-{j}';sample=AuditSample(sid,sid,pd.DataFrame({'x':effect+rng.normal()+rng.normal(size=particles)}));s.append(sample);mapping[sid]=str(b)
    return s,mapping

def test_hierarchy_recovers_known_variance_components():
    s,m=hierarchical();r=estimate(sufficient_stats(s,'x',m))
    assert r['model_status']=='converged'
    assert r['sigma2_bio']==pytest.approx(4.,abs=2.)
    assert r['sigma2_nested_sample']==pytest.approx(1.,abs=.5)
    assert r['sigma2_particle']==pytest.approx(1.,abs=.08)
    assert r['icc_same_nested_sample']==pytest.approx((r['sigma2_bio']+r['sigma2_nested_sample'])/sum(r[k] for k in ['sigma2_bio','sigma2_nested_sample','sigma2_particle']))

def test_reml_matches_full_covariance_independent_likelihood():
    from scipy.optimize import minimize
    s,m=hierarchical(seed=14,bios=4,nested=3,particles=6)
    y=np.concatenate([x.dataframe.x.to_numpy() for x in s]);n=len(y);bi=np.repeat([m[x.sample_id] for x in s],6);si=np.repeat(np.arange(len(s)),6)
    B=(bi[:,None]==bi).astype(float);S=(si[:,None]==si).astype(float);one=np.ones(n)
    def objective(theta):
        V=theta[0]*B+theta[1]*S+theta[2]*np.eye(n)
        sign,logdet=np.linalg.slogdet(V)
        if sign<=0:return 1e20
        inv=np.linalg.solve(V,np.column_stack([y,one]));xx=one@inv[:,1];xy=one@inv[:,0]
        return .5*(logdet+np.log(xx)+y@inv[:,0]-xy*xy/xx)
    full=minimize(objective,[2,1,1],method='L-BFGS-B',bounds=[(0,None),(0,None),(1e-10,None)])
    optimized=estimate(sufficient_stats(s,'x',m));actual=[optimized[k] for k in ('sigma2_bio','sigma2_nested_sample','sigma2_particle')]
    assert np.allclose(actual,full.x,rtol=2e-3,atol=2e-3)

def test_hierarchy_unbalanced_equal_bio_influence_and_identifiability():
    s,m=hierarchical(bios=5,nested=3,particles=30)
    s=[sample for sample in s if sample.sample_id!='0-2']
    m={sample.sample_id:m[sample.sample_id] for sample in s}
    for sample in s:
        if m[sample.sample_id]=='4':sample.dataframe.x+=25
        if m[sample.sample_id]=='0':sample.dataframe=pd.concat([sample.dataframe]*10,ignore_index=True)
    r=run_audit(s,audit_type='hierarchy',hierarchy_mapping=m,policy={'n_projections':4})
    assert r['variance_components'][0]['n_biological']==5
    expected=np.mean([np.mean([x.dataframe.x.mean() for x in s if m[x.sample_id]==b]) for b in set(m.values())])
    assert r['equal_bio_centers'][0]['equal_bio_center']==pytest.approx(expected)
    biggest=max(r['influence'],key=lambda x:x['absolute_change']);assert biggest['omitted_bio']=='4' and biggest['influence_status']=='influential'
    reduced=s[::3];mapping={x.sample_id:m[x.sample_id] for x in reduced}
    out=estimate(sufficient_stats(reduced,'x',mapping));assert out['model_status']=='unidentifiable' and out['sigma2_bio'] is None

def test_multiple_audits_migration_immutable_history_and_export_parity(tmp_path):
    s=samples();w=Workspace()
    for sample in s:
        p=tmp_path/(sample.sample_id+'.csv');sample.dataframe.to_csv(p,index=False)
        ref=FileReference.create(p);sample.source=copy.deepcopy(ref.__dict__);w.samples[sample.sample_id]=SampleReference(sample.sample_id,sample.display_name,ref)
    bundle=run_audit(s);ref=commit_bundle(w,bundle,tmp_path/'test.vflow')
    second=commit_bundle(w,run_audit(s,name='another'),tmp_path/'test.vflow')
    baseline=Path(ref.result.absolute_path).read_bytes() if False else open(ref.result.absolute_path,'rb').read()
    append_decision(ref,'0','exclude',note='Review note');append_decision(ref,'0','restore')
    assert open(ref.result.absolute_path,'rb').read()==baseline
    assert read_bundle(ref,tmp_path/'test.vflow')==bundle
    dest=tmp_path/'test.vflow';w.save(dest);reopened=Workspace.load(dest)
    assert len(reopened.audits)==2 and len(reopened.audits[ref.audit_id].decision_history)==2
    xlsx=tmp_path/'audit.xlsx';csv=tmp_path/'audit.zip';export_xlsx(bundle,xlsx,ref,w);export_csv_package(bundle,csv,ref,w)
    from openpyxl import load_workbook
    book=load_workbook(xlsx,data_only=True);assert tuple(book.sheetnames)==SHEETS
    sheet=book['Global_Distance_Matrix'];assert sheet.cell(2,3).value==pytest.approx(bundle['global_distance_matrix'][0][1])
    with zipfile.ZipFile(csv) as z:
        assert json.loads(z.read('canonical_audit.json'))==bundle
        assert len([n for n in z.namelist() if n.endswith('.csv')])==len(SHEETS)
    old=json.loads(dest.read_text());old['workspace_schema']=1;old.pop('audits');dest.write_text(json.dumps(old));migrated=Workspace.load(dest)
    assert migrated.workspace_schema==2 and migrated.audits=={}
    with pytest.raises(ValueError,match='already'):commit_bundle(w,bundle,dest)
    open(ref.result.absolute_path,'wb').write(b'changed')
    with pytest.raises(ValueError,match='missing or changed'):read_bundle(ref,dest)
